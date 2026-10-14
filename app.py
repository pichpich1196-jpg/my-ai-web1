import os
import tempfile
import docx
from google import genai
from pptx import Presentation
import streamlit as st

# កំណត់ Layout ទំព័រ
st.set_page_config(
    page_title="AI Smart Workspace",
    page_icon="⚡",
    layout="wide",
    initial_sidebar_state="expanded",
)

# ==================== MODERN CUSTOM CSS ====================
st.markdown(
    """
    <style>
    @import url('https://fonts.googleapis.com/css2?family=Kantumruy+Pro:wght@300;400;500;600;700&display=swap');

    * {
        font-family: 'Kantumruy Pro', sans-serif !important;
    }

    .hero-container {
        background: linear-gradient(135deg, #1e3c72 0%, #2a5298 100%);
        padding: 30px;
        border-radius: 16px;
        color: white;
        margin-bottom: 25px;
        box-shadow: 0 8px 24px rgba(0, 0, 0, 0.12);
        text-align: center;
    }
    .hero-title {
        font-size: 30px;
        font-weight: 700;
        margin-bottom: 8px;
    }
    .hero-subtitle {
        font-size: 15px;
        opacity: 0.9;
        font-weight: 300;
    }

    div.stButton > button:first-child {
        background: linear-gradient(135deg, #0062E0 0%, #1976D2 100%);
        color: white;
        border-radius: 10px;
        padding: 10px 24px;
        font-size: 16px;
        font-weight: 600;
        border: none;
        width: 100%;
        transition: all 0.3s ease;
    }
    div.stButton > button:first-child:hover {
        transform: translateY(-2px);
        box-shadow: 0 6px 20px rgba(25, 118, 210, 0.4);
    }
    </style>
    """,
    unsafe_allow_html=True,
)

# ទាញយក API Key ពី Streamlit Secrets ប្រកបដោយសុវត្ថិភាព
api_key = st.secrets.get("GEMINI_API_KEY", "")
if not api_key:
  st.warning("⚠️ សូមកំណត់ GEMINI_API_KEY នៅក្នុង Streamlit Secrets ជាមុនសិន!")
  st.stop()

client = genai.Client(api_key=api_key)

# ==================== HEADER ====================
st.markdown(
    """
    <div class="hero-container">
        <div class="hero-title">⚡ ប្រព័ន្ធជំនួយការឆ្លាតវៃ AI Workspace</div>
        <div class="hero-subtitle">ជំនួយការបច្ចេកទេស និងរដ្ឋបាលស្វ័យប្រវត្តិកម្រិតខ្ពស់ (សំឡេង ឯកសារ រូបភាព និងពិគ្រោះយោបល់)</div>
    </div>
    """,
    unsafe_allow_html=True,
)

# ==================== MAIN TABS ====================
tab1, tab2, tab3, tab4 = st.tabs([
    "🎙️ សំឡេង ➜ របាយការណ៍",
    "📄 សម្រាយ & សង្ខេបឯកសារ",
    "🖼️ អានក្រដាស Scan / OCR",
    "💬 AI ពិគ្រោះ & ព្រាងលិខិត",
])

# ----------------- TAB 1: VOICE REPORT -----------------
with tab1:
  col_left, col_right = st.columns([1, 1], gap="medium")
  with col_left:
    st.markdown("### 🎙️ បញ្ចូលសំឡេង")
    audio_file = st.file_uploader(
        "ជ្រើសរើស File សំឡេង (MP3, WAV, OGG, M4A)៖",
        type=["mp3", "wav", "ogg", "m4a"],
        key="aud_up",
    )
    report_type = st.selectbox(
        "ជ្រើសរើសទម្រង់របាយការណ៍ដែលត្រូវការ៖",
        [
            "📊 របាយការណ៍វឌ្ឍនភាពការងារទូទៅ",
            "🏛️ របាយការណ៍ចុះមូលដ្ឋាន",
            "🛠️ របាយការណ៍បច្ចេកទេស & ដោះស្រាយបញ្ហា",
            "📝 កំណត់ហេតុអង្គប្រជុំ (Minutes of Meeting)",
        ],
    )
    btn_audio = st.button("✨ រៀបចំរបាយការណ៍ដោយ AI", key="btn_run_audio")

  with col_right:
    st.markdown("### 📋 លទ្ធផលរបាយការណ៍")
    if audio_file and btn_audio:
      with st.spinner("AI កំពុងស្ដាប់ និងរៀបចំជារបាយការណ៍ផ្លូវការ..."):
        ext = audio_file.name.split(".")[-1]
        with tempfile.NamedTemporaryFile(
            delete=False, suffix=f".{ext}"
        ) as tmp:
          tmp.write(audio_file.read())
          tmp_path = tmp.name

        try:
          uploaded = client.files.upload(
              file=tmp_path, mime_type=audio_file.type
          )
          prompt = f"""
                    សូមស្ដាប់សំឡេងនេះ ហើយរៀបចំចេញជា {report_type} ជាភាសាខ្មែរផ្លូវការ ស្អាតបាត។
                    បែងចែកជាចំណុចច្បាស់ៗ៖
                    ១. ព័ត៌មានទូទៅ (កាលបរិច្ឆេទ ទីតាំង សមាសភាព)
                    ២. សកម្មភាពស្នូល
                    ៣. លទ្ធផលសម្រេចបាន
                    ៤. បញ្ហាប្រឈម & ដំណោះស្រាយ
                    ៥. ផែនការអនុវត្តបន្ត
                    """
          resp = client.models.generate_content(
              model="gemini-2.5-flash", contents=[uploaded, prompt]
          )
          client.files.delete(name=uploaded.name)
          st.success("✅ រៀបចំរបាយការណ៍រួចរាល់!")
          st.markdown(resp.text)
        except Exception as e:
          st.error(f"កំហុស៖ {e}")
        finally:
          if os.path.exists(tmp_path):
            os.remove(tmp_path)
    else:
      st.info("👈 សូមជ្រើសរើស File សំឡេងនៅខាងឆ្វេង រួចចុចប៊ូតុងដើម្បីមើលលទ្ធផល")

# ----------------- TAB 2: DOCUMENT SUMMARY -----------------
with tab2:
  col_doc1, col_doc2 = st.columns([1, 1], gap="medium")
  with col_doc1:
    st.markdown("### 📄 ជ្រើសរើសឯកសារ")
    doc_file = st.file_uploader(
        "បញ្ចូល File (PDF, Word, PowerPoint)៖",
        type=["pdf", "docx", "pptx"],
        key="doc_up",
    )
    doc_mode = st.radio(
        "កម្រិតនៃការសម្រាយ៖",
        ["សង្ខេបចំណុចស្នូល ៣ ទៅ ៥ ចំណុច", "សម្រាយខ្លឹមសារលម្អិតតាមជំពូក"],
        horizontal=True,
    )
    btn_doc = st.button("✨ សម្រាយឯកសារឥឡូវនេះ", key="btn_run_doc")

  with col_doc2:
    st.markdown("### 📑 ខ្លឹមសារសង្ខេប")
    if doc_file and btn_doc:
      with st.spinner("AI កំពុងអាន និងសម្រាយខ្លឹមសារ..."):
        ext = doc_file.name.split(".")[-1].lower()
        with tempfile.NamedTemporaryFile(delete=False, suffix=f".{ext}") as tmp:
          tmp.write(doc_file.read())
          tmp_path = tmp.name

        try:
          p_doc = (
              "សូមសម្រាយខ្លឹមសារនេះជាភាសាខ្មែរតាមទម្រង់ '"
              + str(doc_mode)
              + "' ដោយប្រើភាសាបច្ចេកទេសត្រឹមត្រូវ និងងាយយល់។"
          )
          if ext == "pdf":
            up = client.files.upload(file=tmp_path, mime_type="application/pdf")
            res = client.models.generate_content(
                model="gemini-2.5-flash", contents=[up, p_doc]
            )
            client.files.delete(name=up.name)
            st.markdown(res.text)
          elif ext == "docx":
            d = docx.Document(tmp_path)
            txt = "\n".join([p.text for p in d.paragraphs if p.text.strip()])
            res = client.models.generate_content(
                model="gemini-2.5-flash", contents=f"{p_doc}\n\n{txt}"
            )
            st.markdown(res.text)
          elif ext == "pptx":
            prs = Presentation(tmp_path)
            txt = "\n".join([
                p.text
                for s in prs.slides
                for sh in s.shapes
                if sh.has_text_frame
                for p in sh.text_frame.paragraphs
                if p.text.strip()
            ])
            res = client.models.generate_content(
                model="gemini-2.5-flash", contents=f"{p_doc}\n\n{txt}"
            )
            st.markdown(res.text)
        except Exception as e:
          st.error(f"កំហុស៖ {e}")
        finally:
          if os.path.exists(tmp_path):
            os.remove(tmp_path)
    else:
      st.info("👈 បញ្ចូលឯកសារដើម្បីឱ្យ AI ជួយសម្រាយជាចំណុចៗ")

# ----------------- TAB 3: IMAGE OCR -----------------
with tab3:
  col_img1, col_img2 = st.columns([1, 1], gap="medium")
  with col_img1:
    st.markdown("### 🖼️ រូបភាព ឬក្រដាស Scan")
    img_file = st.file_uploader(
        "បញ្ចូលរូបភាព (JPG, PNG)៖", type=["jpg", "jpeg", "png"], key="img_up"
    )
    img_action = st.selectbox("តម្រូវការអនុវត្ត៖", [
        "ស្រង់អក្សរចេញទាំងអស់ (Extract Text / OCR)",
        "បកប្រែខ្លឹមសារក្នុងរូបភាពជាភាសាខ្មែរ",
        "វិភាគទិន្នន័យ ឬតារាងក្នុងរូបភាព",
    ])
    btn_img = st.button("✨ ដំណើរការរូបភាព", key="btn_run_img")

  with col_img2:
    st.markdown("### 📝 លទ្ធផលស្រង់ចេញ")
    if img_file and btn_img:
      with st.spinner("AI កំពុងវិភាគរូបភាព..."):
        ext = img_file.name.split(".")[-1]
        with tempfile.NamedTemporaryFile(delete=False, suffix=f".{ext}") as tmp:
          tmp.write(img_file.read())
          tmp_path = tmp.name

        try:
          up = client.files.upload(file=tmp_path, mime_type=img_file.type)
          res = client.models.generate_content(
              model="gemini-2.5-flash",
              contents=[up, f"ធ្វើសកម្មភាព៖ {img_action} ជាភាសាខ្មែរ"],
          )
          client.files.delete(name=up.name)
          st.markdown(res.text)
        except Exception as e:
          st.error(f"កំហុស៖ {e}")
        finally:
          if os.path.exists(tmp_path):
            os.remove(tmp_path)
    else:
      st.info("👈 បញ្ចូលរូបភាព Scan ឬឯកសារថត")

# ----------------- TAB 4: AI CONSULTANT -----------------
with tab4:
  st.markdown("### 💬 ជំនួយការព្រាងលិខិត & ពិគ្រោះយោបល់")
  user_query = st.text_area(
      "សរសេរប្រធានបទ ឬសំណួររបស់អ្នកនៅទីនេះ (ឧ. សេចក្តីជូនដំណឹង, លិខិតផ្លូវការ, ដំណោះស្រាយបច្ចេកទេស)៖",
      height=140,
  )
  if st.button("✨ បង្កើតខ្លឹមសារ", key="btn_run_chat") and user_query:
    with st.spinner("AI កំពុងរៀបចំខ្លឹមសារ..."):
      res = client.models.generate_content(
          model="gemini-2.5-flash",
          contents=(
              "អ្នកគឺជាជំនួយការបច្ចេកទេស និងរដ្ឋបាលជាន់ខ្ពស់។"
              f" សូមឆ្លើយតបជាភាសាខ្មែរផ្លូវការ សុភាព និងមានលំដាប់លំដោយ៖ {user_query}"
          ),
      )
      st.markdown("---")
      st.markdown(res.text)
